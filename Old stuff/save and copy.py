import os
import shutil
import sqlite3
from fpdf import FPDF
from docx import Document
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.filechooser import FileChooserIconView
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.popup import Popup

# Database setup
def init_db():
    conn = sqlite3.connect('file_operations.db')
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS file_operations (
                        id INTEGER PRIMARY KEY,
                        filename TEXT,
                        operation TEXT)''')
    conn.commit()
    conn.close()

def log_operation(filename, operation):
    conn = sqlite3.connect('file_operations.db')
    cursor = conn.cursor()
    cursor.execute("INSERT INTO file_operations (filename, operation) VALUES (?, ?)", (filename, operation))
    conn.commit()
    conn.close()

# File conversion functions
def convert_txt_to_pdf(txt_file, pdf_file):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    with open(txt_file, "r") as file:
        for line in file:
            pdf.cell(200, 10, txt=line, ln=True)

    pdf.output(pdf_file)
    log_operation(txt_file, "Converted to PDF")

def convert_docx_to_txt(docx_file, txt_file):
    doc = Document(docx_file)
    text = "\n".join([para.text for para in doc.paragraphs])

    with open(txt_file, "w") as file:
        file.write(text)
    log_operation(docx_file, "Converted to Text")

def copy_file(src, dest):
    shutil.copy(src, dest)
    log_operation(src, f"Copied to {dest}")

# Kivy App UI
class FileManagerApp(App):
    def build(self):
        self.layout = BoxLayout(orientation='vertical')

        # File chooser
        self.filechooser = FileChooserIconView()
        self.layout.add_widget(self.filechooser)

        # Convert buttons
        self.txt_to_pdf_button = Button(text="Convert TXT to PDF")
        self.txt_to_pdf_button.bind(on_press=self.convert_txt_to_pdf)
        self.layout.add_widget(self.txt_to_pdf_button)

        self.docx_to_txt_button = Button(text="Convert DOCX to TXT")
        self.docx_to_txt_button.bind(on_press=self.convert_docx_to_txt)
        self.layout.add_widget(self.docx_to_txt_button)

        # Copy file button
        self.copy_button = Button(text="Copy File")
        self.copy_button.bind(on_press=self.copy_file)
        self.layout.add_widget(self.copy_button)

        # Log output
        self.log_output = TextInput(size_hint=(1, 0.2), readonly=True)
        self.layout.add_widget(self.log_output)

        return self.layout

    def convert_txt_to_pdf(self, instance):
        if self.filechooser.selection:
            txt_file = self.filechooser.selection[0]
            pdf_file = os.path.splitext(txt_file)[0] + '.pdf'
            convert_txt_to_pdf(txt_file, pdf_file)
            self.show_popup(f"Converted {txt_file} to {pdf_file}")

    def convert_docx_to_txt(self, instance):
        if self.filechooser.selection:
            docx_file = self.filechooser.selection[0]
            txt_file = os.path.splitext(docx_file)[0] + '.txt'
            convert_docx_to_txt(docx_file, txt_file)
            self.show_popup(f"Converted {docx_file} to {txt_file}")

    def copy_file(self, instance):
        if self.filechooser.selection:
            src = self.filechooser.selection[0]
            dest = os.path.join(os.path.dirname(src), "copy_" + os.path.basename(src))
            copy_file(src, dest)
            self.show_popup(f"Copied {src} to {dest}")

    def show_popup(self, message):
        popup = Popup(title='Operation Result',
                      content=Label(text=message),
                      size_hint=(0.6, 0.6))
        popup.open()

if __name__ == '__main__':
    init_db()
    FileManagerApp().run()
